import os
import fitz
import tempfile

_ocr_enabled = True

try:
    import pytesseract
    from PIL import Image
    pytesseract.get_tesseract_version()
except Exception:
    _ocr_enabled = False


def _ocr_page(page):
    if not _ocr_enabled:
        return ""
    pix = page.get_pixmap(dpi=200)
    img_data = pix.tobytes("png")
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
        tmp.write(img_data)
        tmp_path = tmp.name
    try:
        text = pytesseract.image_to_string(Image.open(tmp_path), lang="eng")
    except Exception:
        text = ""
    finally:
        os.unlink(tmp_path)
    return text


def extract_text(file_path):

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        doc = fitz.open(file_path)
        if doc.is_encrypted:
            auth = doc.authenticate("")
            if auth == 0:
                doc.close()
                raise ValueError("PDF is password-protected. Please provide an unlocked PDF.")

        all_text = ""

        for i in range(doc.page_count):
            page = doc.load_page(i)
            page_text = page.get_text("text")
            if page_text.strip():
                all_text += page_text + "\n"
            elif _ocr_enabled:
                ocr_text = _ocr_page(page)
                if ocr_text.strip():
                    all_text += ocr_text + "\n"

        doc.close()
        return all_text

    if ext == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    if ext == ".docx":
        try:
            from docx import Document
            doc = Document(file_path)
            return "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
        except ImportError:
            raise ImportError("python-docx required. pip install python-docx")

    if ext in (".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".tif", ".webp"):
        if not _ocr_enabled:
            raise ValueError("OCR not available. Install tesseract-ocr and pytesseract.")
        return pytesseract.image_to_string(Image.open(file_path), lang="eng")

    if ext == ".csv":
        import csv
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            return "\n".join([", ".join(row) for row in reader])

    if ext == ".json":
        import json
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            data = json.load(f)
            return json.dumps(data, indent=2)

    if ext == ".xlsx":
        try:
            from openpyxl import load_workbook
            wb = load_workbook(file_path, read_only=True, data_only=True)
            all_text = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                all_text.append(f"=== Sheet: {sheet} ===")
                for row in ws.iter_rows(values_only=True):
                    row_text = ", ".join([str(cell) if cell is not None else "" for cell in row])
                    if row_text.strip(", "):
                        all_text.append(row_text)
            wb.close()
            return "\n".join(all_text)
        except ImportError:
            raise ImportError("openpyxl required. pip install openpyxl")

    if ext == ".pptx":
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            all_text = []
            for i, slide in enumerate(prs.slides, 1):
                slide_text = [f"=== Slide {i} ==="]
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            text = para.text.strip()
                            if text:
                                slide_text.append(text)
                    if shape.has_table:
                        for row in shape.table.rows:
                            row_text = ", ".join([cell.text.strip() for cell in row.cells])
                            if row_text.strip(", "):
                                slide_text.append(row_text)
                all_text.append("\n".join(slide_text))
            return "\n\n".join(all_text)
        except ImportError:
            raise ImportError("python-pptx required. pip install python-pptx")

    if ext == ".md":
        try:
            import re
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            content = re.sub(r'```[\s\S]*?```', lambda m: m.group(0), content)
            content = re.sub(r'[#*_`~>\-|]', ' ', content)
            content = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', content)
            content = re.sub(r'!\[([^\]]*)\]\([^\)]+\)', r'[Image: \1]', content)
            content = re.sub(r'\n{3,}', '\n\n', content)
            return content.strip()
        except Exception as e:
            raise ValueError(f"Error reading markdown: {e}")

    if ext == ".html":
        try:
            import re
            from bs4 import BeautifulSoup
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                soup = BeautifulSoup(f.read(), "html.parser")
            for tag in soup(["script", "style", "nav", "footer", "header"]):
                tag.decompose()
            text = soup.get_text(separator="\n", strip=True)
            text = re.sub(r'\n{3,}', '\n\n', text)
            return text.strip()
        except ImportError:
            raise ImportError("beautifulsoup4 required. pip install beautifulsoup4")

    raise ValueError(f"Unsupported file format: {ext}. Supported: PDF, TXT, DOCX, PNG, JPG, JPEG, BMP, TIFF, WEBP, CSV, JSON, XLSX, PPTX, MD, HTML")